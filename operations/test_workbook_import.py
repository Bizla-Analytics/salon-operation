import gzip
import io
import json
from datetime import timedelta
from unittest.mock import patch
from zipfile import ZipFile

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.files.uploadhandler import StopUpload
from django.test import Client, TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook

from .models import Branch, Customer, OperationalTask, Service, SOPWorkbookImport, Visit, VisitService, VisitTask
from .sop_imports import MAX_UPLOAD, SCHEMA, ImportProblem, validate_workbook
from .upload_limits import WorkbookUploadLimit
from .workflow import build_visit_tasks


def synthetic_workbook(change=None):
    records = {
        "Service Master": [{"service_code": "TEST_SOP", "service_name": "Synthetic service", "standard_duration_minutes": 10}],
        "Sub-Service Master": [
            {"sub_service_code": code, "sub_service_name": name}
            for code, name in [("SUB025", "Sanitisation"), ("SUB001", "Consultation"), ("TEST_SUB", "Synthetic procedure")]
        ],
        "Task Master": [{"task_id": "TEST_TASK_" + code, "sub_service_code": code, "task_name": name,
                         "task_sequence": 1, "active_labour_minutes": 2, "is_skippable": False}
                        for code, name in [("SUB025", "Sanitise"), ("SUB001", "Consult"), ("TEST_SUB", "Synthetic task")]],
        "Service Detail": [{"service_detail_id": "TEST_DETAIL", "service_code": "TEST_SOP", "sub_service_code": "TEST_SUB",
                            "sequence": 1, "required_quantity": 1, "mandatory": True}],
    }
    if change:
        change(records)
    book = Workbook()
    book.remove(book.active)
    for name, schema in SCHEMA.items():
        headers = schema.split()
        sheet = book.create_sheet(name)
        sheet.append(headers)
        for record in records.get(name, []):
            # Real Excel omits trailing empty cells; exercise that format as well.
            sheet.append([record.get(header) for header in headers])
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


class WorkbookImportTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user("sop_admin", password="test-password")
        self.admin.profile.role = "ADMIN"
        self.admin.profile.save()
        self.client.force_login(self.admin)
        self.url = reverse("workbook_import")

    def upload(self, payload=None, filename="synthetic.xlsx"):
        return self.client.post(self.url, {
            "action": "validate", "workbook": SimpleUploadedFile(filename, payload or synthetic_workbook()),
        })

    def confirm(self, stage):
        return self.client.post(self.url, {"action": "confirm", "stage": str(stage.pk), "confirmed": "on"})

    def test_anonymous_redirects_to_login(self):
        self.client.logout()
        self.assertEqual(self.client.get(self.url).status_code, 302)
        self.assertEqual(self.upload().status_code, 302)

    def test_branch_roles_cannot_import_or_download(self):
        for role in ["GENERAL_MANAGER", "MANAGER", "EMPLOYEE"]:
            self.admin.profile.role = role
            self.admin.profile.save()
            self.assertEqual(self.client.get(self.url).status_code, 403)
            self.assertEqual(self.upload().status_code, 403)
            self.assertEqual(self.client.get(reverse("workbook_snapshot", args=["00000000-0000-0000-0000-000000000001"])).status_code, 403)

    def test_inactive_admin_denied(self):
        self.admin.profile.active = False
        self.admin.profile.save()
        self.assertEqual(self.upload().status_code, 403)

    def test_superuser_allowed(self):
        self.admin.is_superuser = True
        self.admin.save()
        self.admin.profile.role = "EMPLOYEE"
        self.admin.profile.save()
        self.assertEqual(self.client.get(self.url).status_code, 200)

    def test_csrf_is_required(self):
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.admin)
        self.assertEqual(client.post(self.url, {"action": "validate"}).status_code, 403)

    def test_preview_rolls_back_and_is_private(self):
        self.assertEqual(self.upload().status_code, 302)
        self.assertEqual(Service.objects.count(), 0)
        self.assertEqual(OperationalTask.objects.count(), 0)
        stage = SOPWorkbookImport.objects.get()
        self.assertEqual(stage.status, "PENDING")
        self.assertTrue(bytes(stage.workbook))
        response = self.client.get(self.url, {"stage": str(stage.pk)})
        self.assertContains(response, "Confirm import")
        self.assertIn("no-store", response["Cache-Control"])
        self.assertContains(response, "total records after import")

    def test_import_and_snapshot_download(self):
        existing = Service.objects.create(code="EXISTING", name="Keep this record", base_price=20)
        self.upload()
        stage = SOPWorkbookImport.objects.get()
        self.assertEqual(self.confirm(stage).status_code, 302)
        stage.refresh_from_db()
        self.assertEqual(stage.status, "APPLIED")
        self.assertEqual(bytes(stage.workbook), b"")
        self.assertEqual(Service.objects.count(), 2)
        self.assertEqual(OperationalTask.objects.count(), 3)
        response = self.client.get(reverse("workbook_snapshot", args=[stage.pk]))
        self.assertEqual(response.status_code, 200)
        fixture = json.loads(gzip.decompress(response.content))
        self.assertEqual(fixture[0]["pk"], existing.pk)
        self.assertEqual(fixture[0]["fields"]["base_price"], "20.00")
        self.assertIn("no-store", response["Cache-Control"])

    def test_confirm_requires_checkbox(self):
        self.upload()
        stage = SOPWorkbookImport.objects.get()
        response = self.client.post(self.url, {"action": "confirm", "stage": str(stage.pk)})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(Service.objects.count(), 0)

    def test_repeat_confirm_is_not_applied_again(self):
        self.upload()
        stage = SOPWorkbookImport.objects.get()
        self.confirm(stage)
        self.confirm(stage)
        self.assertEqual(Service.objects.count(), 1)
        self.assertEqual(OperationalTask.objects.count(), 3)

    def test_other_admin_cannot_access_confirm_cancel_or_download(self):
        self.upload()
        stage = SOPWorkbookImport.objects.get()
        other = User.objects.create_superuser("other_sop_admin", password="test-password")
        self.client.force_login(other)
        self.assertEqual(self.client.get(self.url, {"stage": str(stage.pk)}).status_code, 404)
        self.assertEqual(self.confirm(stage).status_code, 404)
        self.assertEqual(self.client.post(self.url, {"action": "cancel", "stage": str(stage.pk)}).status_code, 404)
        self.assertEqual(self.client.get(reverse("workbook_snapshot", args=[stage.pk])).status_code, 404)
        self.assertEqual(Service.objects.count(), 0)

    def test_cancel_removes_staged_file(self):
        self.upload()
        stage = SOPWorkbookImport.objects.get()
        self.client.post(self.url, {"action": "cancel", "stage": str(stage.pk)})
        stage.refresh_from_db()
        self.assertEqual(stage.status, "CANCELLED")
        self.assertEqual(bytes(stage.workbook), b"")
        self.assertEqual(Service.objects.count(), 0)

    def test_new_upload_discards_previous_preview(self):
        self.upload()
        first = SOPWorkbookImport.objects.get()
        self.upload()
        first.refresh_from_db()
        self.assertEqual(first.status, "CANCELLED")
        self.assertEqual(bytes(first.workbook), b"")
        self.assertEqual(SOPWorkbookImport.objects.filter(status="PENDING").count(), 1)

    def test_expiry_rejects_and_clears_file(self):
        self.upload()
        stage = SOPWorkbookImport.objects.get()
        stage.expires_at = timezone.now() - timedelta(seconds=1)
        stage.save()
        self.confirm(stage)
        stage.refresh_from_db()
        self.assertEqual(stage.status, "EXPIRED")
        self.assertEqual(bytes(stage.workbook), b"")
        self.assertEqual(Service.objects.count(), 0)

    def test_master_edit_invalidates_preview(self):
        self.upload()
        stage = SOPWorkbookImport.objects.get()
        Service.objects.create(code="EDITED", name="Another administrator's change")
        response = self.confirm(stage)
        self.assertContains(response, "Master data changed")
        self.assertFalse(Service.objects.filter(code="TEST_SOP").exists())

    def test_wrong_extension_and_broken_file_rejected(self):
        self.assertContains(self.upload(filename="bad.xlsm"), "Choose an .xlsx")
        self.assertContains(self.upload(payload=b"not an xlsx"), "not a readable")
        self.assertEqual(SOPWorkbookImport.objects.count(), 0)

    def test_large_upload_rejected_before_file_parsing(self):
        response = self.client.post(self.url, data=b"x" * (MAX_UPLOAD + 1024 * 1024 + 1), content_type="application/octet-stream")
        self.assertEqual(response.status_code, 413)

    def test_chunked_file_size_limit(self):
        handler = WorkbookUploadLimit()
        self.assertEqual(handler.receive_data_chunk(b"test", 0), b"test")
        with self.assertRaises(StopUpload):
            handler.receive_data_chunk(b"test", MAX_UPLOAD)

    def test_macro_content_rejected_even_with_xlsx_extension(self):
        output = io.BytesIO(synthetic_workbook())
        with ZipFile(output, "a") as archive:
            archive.writestr("xl/vbaProject.bin", b"synthetic macro placeholder")
        self.assertContains(self.upload(payload=output.getvalue()), "Macro-enabled")

    def test_missing_sheet_rejected(self):
        book = Workbook()
        output = io.BytesIO()
        book.save(output)
        book.close()
        self.assertContains(self.upload(payload=output.getvalue()), "Missing sheets")

    def test_row_limit_rejects_before_import(self):
        with patch("operations.sop_imports.MAX_ROWS", 1):
            with self.assertRaisesMessage(ImportProblem, "exceeds 10,000 rows"):
                validate_workbook(synthetic_workbook())

    def test_duplicate_codes_rejected(self):
        payload = synthetic_workbook(lambda rows: rows["Service Master"].append(dict(rows["Service Master"][0])))
        self.assertContains(self.upload(payload=payload), "duplicate identifier")
        self.assertEqual(Service.objects.count(), 0)

    def test_invalid_numeric_values_do_not_silently_become_zero(self):
        for value in ["invalid", "NaN", "Infinity", -1, 2.5]:
            payload = synthetic_workbook(lambda rows: rows["Task Master"][0].update(active_labour_minutes=value))
            self.assertContains(self.upload(payload=payload), "valid non-negative number")

    def test_invalid_reference_rolls_back_whole_preview(self):
        payload = synthetic_workbook(lambda rows: rows["Task Master"][0].update(sub_service_code="NOT_FOUND"))
        self.assertContains(self.upload(payload=payload), "validation failed")
        self.assertEqual(Service.objects.count(), 0)
        self.assertEqual(SOPWorkbookImport.objects.count(), 0)

    def test_required_opening_tasks_validated(self):
        payload = synthetic_workbook(lambda rows: rows["Task Master"].pop(0))
        self.assertContains(self.upload(payload=payload), "Workflow requires SUB025")
        self.assertEqual(Service.objects.count(), 0)

    def test_failed_import_rolls_back_snapshot_and_all_changes(self):
        self.upload()
        stage = SOPWorkbookImport.objects.get()

        def fail_after_write(*args, **kwargs):
            Service.objects.create(code="ROLLBACK", name="Must not persist")
            raise ImportProblem("Synthetic failure")

        with patch("operations.sop_imports.run_import", side_effect=fail_after_write):
            self.assertContains(self.confirm(stage), "Synthetic failure")
        self.assertFalse(Service.objects.exists())
        stage.refresh_from_db()
        self.assertEqual(stage.status, "PENDING")
        self.assertEqual(bytes(stage.master_snapshot), b"")

    def test_staged_workbook_tampering_rejected(self):
        self.upload()
        stage = SOPWorkbookImport.objects.get()
        SOPWorkbookImport.objects.filter(pk=stage.pk).update(workbook=b"tampered")
        self.assertContains(self.confirm(stage), "staged workbook changed")
        self.assertEqual(Service.objects.count(), 0)

    def test_existing_visit_snapshots_not_rewritten_by_reimport(self):
        self.upload()
        self.confirm(SOPWorkbookImport.objects.get())
        branch = Branch.objects.create(code="TEST_BRANCH", name="Synthetic branch")
        customer = Customer.objects.create(name="Synthetic customer")
        visit = Visit.objects.create(branch=branch, customer=customer, created_by=self.admin, status="ASSIGNED")
        VisitService.objects.create(visit=visit, service=Service.objects.get(code="TEST_SOP"),
                                    assigned_by=self.admin, employee=self.admin, order_number=1)
        build_visit_tasks(visit)
        before = list(VisitTask.objects.values())
        payload = synthetic_workbook(lambda rows: rows["Task Master"][2].update(task_name="Changed master task"))
        self.upload(payload=payload)
        self.confirm(SOPWorkbookImport.objects.filter(status="PENDING").get())
        self.assertEqual(list(VisitTask.objects.values()), before)
        self.assertTrue(OperationalTask.objects.filter(name="Changed master task").exists())

    def test_invalid_uuid_and_unknown_action_rejected(self):
        self.assertEqual(self.client.get(self.url, {"stage": "invalid"}).status_code, 400)
        self.assertEqual(self.client.post(self.url, {"action": "unknown"}).status_code, 400)

    def test_new_routes_visible_in_admin_ui(self):
        self.assertContains(self.client.get(reverse("admin_dashboard")), reverse("import_data"))
        self.assertContains(self.client.get(reverse("import_data")), self.url)
        self.assertRedirects(self.client.get(reverse("csv_import")), reverse("import_data"))
        self.assertContains(self.client.get(reverse("import_data")), "Excel SOP workbook")
