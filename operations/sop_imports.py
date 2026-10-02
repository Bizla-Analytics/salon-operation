"""Bounded uploads using the same atomic importer as the CLI; no public media."""
import gzip
import hashlib
import io
import tempfile
from contextlib import contextmanager
from datetime import timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from django.core import serializers
from django.core.exceptions import ObjectDoesNotExist
from django.core.management import call_command
from django.core.management.base import CommandError
from django.db import DatabaseError, connection, transaction
from django.utils import timezone
from openpyxl import load_workbook

from .models import (
    Equipment, InventoryItem, OperationalTask, Service, ServiceDetail,
    SOPWorkbookImport, SubService, TaskEquipment, TaskInventory,
)

MAX_UPLOAD = 10 * 1024 * 1024
MAX_EXPANDED = 64 * 1024 * 1024
MAX_ROWS = 10000
MASTER_MODELS = (Service, SubService, InventoryItem, Equipment, OperationalTask,
                 ServiceDetail, TaskInventory, TaskEquipment)
# Required by import_sop_workbook. Extra columns are allowed; missing ones are not.
SCHEMA = {
    "Service Master": "service_code service_name service_category customer_type service_type standard_duration_minutes current_version notes",
    "Sub-Service Master": "sub_service_code sub_service_name category standard_output_quantity output_unit procedure_notes",
    "Inventory Master": "inventory_code inventory_name inventory_category source_method base_unit purchase_pack_quantity purchase_pack_unit resource_type notes",
    "Equipment Master": "equipment_code equipment_name equipment_category utility_type branch_code rated_power_kw water_litres_per_minute notes",
    "Task Master": "task_id sub_service_code task_sequence activity_code task_name role_code employee_role active_labour_minutes passive_time_minutes required_skill allowed_staff_type default_customer_type is_skippable notes",
    "Service Detail": "service_detail_id service_code sub_service_code sequence required_quantity mandatory mapping_status notes",
    "Task Inventory": "task_inventory_id task_id service_code inventory_code standard_quantity unit waste_percent effective_quantity mapping_status notes mapping_note",
    "Task Equipment": "task_equipment_id task_id equipment_code quantity_required equipment_usage_minutes utility_type utility_minutes notes",
}
INTEGER_COLUMNS = {"standard_duration_minutes", "task_sequence", "sequence", "active_labour_minutes",
                   "passive_time_minutes", "equipment_usage_minutes", "utility_minutes"}
DECIMAL_COLUMNS = {"standard_output_quantity", "purchase_pack_quantity", "rated_power_kw",
                   "water_litres_per_minute", "required_quantity", "standard_quantity",
                   "waste_percent", "effective_quantity", "quantity_required"}


class ImportProblem(Exception):
    pass


def validate_workbook(payload):
    if not payload or len(payload) > MAX_UPLOAD:
        raise ImportProblem("Choose an .xlsx workbook no larger than 10 MB.")
    try:
        with ZipFile(io.BytesIO(payload)) as archive:
            entries = archive.infolist()
            if (len(entries) > 1000 or sum(item.file_size for item in entries) > MAX_EXPANDED
                    or any(item.flag_bits & 1 for item in entries)):
                raise ImportProblem("Workbook is encrypted or exceeds the safe workbook limits.")
            if any(item.filename.lower().endswith("vbaproject.bin") for item in entries):
                raise ImportProblem("Macro-enabled workbooks are not supported. Use .xlsx.")
        book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except (BadZipFile, ValueError, KeyError, OSError) as exc:
        raise ImportProblem("This is not a readable .xlsx workbook.") from exc
    try:
        missing = set(SCHEMA).difference(book.sheetnames)
        if missing:
            raise ImportProblem("Missing sheets: " + ", ".join(sorted(missing)))
        total = 0
        for name, schema in SCHEMA.items():
            sheet = book[name]
            # Untrusted dimensions must not hide data or force billions of rows.
            sheet.reset_dimensions()
            iterator = sheet.iter_rows(values_only=True)
            first = next(iterator, ())
            headers = [str(value).strip() if value is not None else "" for value in first]
            meaningful = [header for header in headers if header]
            if len(meaningful) != len(set(meaningful)):
                raise ImportProblem(f"{name}: duplicate column names.")
            absent = set(schema.split()).difference(headers)
            if absent:
                raise ImportProblem(f"{name}: missing columns: {', '.join(sorted(absent))}")
            key = headers.index(schema.split()[0])
            identifiers = set()
            for number, row in enumerate(iterator, 2):
                total += 1
                if total > MAX_ROWS:
                    raise ImportProblem("Workbook exceeds 10,000 rows. Ask an administrator to use the CLI.")
                if not any(value not in (None, "") for value in row):
                    continue
                identifier = str(row[key]).strip() if key < len(row) and row[key] is not None else ""
                if not identifier or identifier in identifiers:
                    raise ImportProblem(f"{name}, row {number}: missing or duplicate identifier.")
                identifiers.add(identifier)
                for column, value in zip(headers, row):
                    if column in INTEGER_COLUMNS | DECIMAL_COLUMNS and value not in (None, ""):
                        try:
                            numeric = Decimal(str(value))
                            valid = numeric.is_finite() and numeric >= 0
                            if column in INTEGER_COLUMNS:
                                valid = valid and numeric == numeric.to_integral_value()
                        except (InvalidOperation, ValueError):
                            valid = False
                        if not valid:
                            raise ImportProblem(f"{name}, row {number}: {column} must be a valid non-negative number.")
            if name in {"Service Master", "Sub-Service Master", "Task Master", "Service Detail"} and not identifiers:
                raise ImportProblem(f"{name}: at least one data row is required.")
    finally:
        book.close()


def expire_stages():
    SOPWorkbookImport.objects.filter(status="PENDING", expires_at__lte=timezone.now()).update(
        status="EXPIRED", workbook=b"",
    )


def master_snapshot():
    objects = [obj for model in MASTER_MODELS for obj in model.objects.order_by("pk")]
    return serializers.serialize("json", objects).encode("utf-8")


def lock_masters():
    # PostgreSQL is authoritative. Also blocks admin/CSV master edits during an import.
    with connection.cursor() as cursor:
        cursor.execute("SELECT current_setting('lock_timeout')")
        previous = cursor.fetchone()[0]
        cursor.execute("SELECT set_config('lock_timeout', '5s', true)")
        tables = sorted(connection.ops.quote_name(model._meta.db_table) for model in MASTER_MODELS)
        cursor.execute("LOCK TABLE " + ", ".join(tables) + " IN SHARE ROW EXCLUSIVE MODE")
        cursor.execute("SELECT set_config('lock_timeout', %s, true)", [previous])


@contextmanager
def private_workbook(payload):
    with tempfile.TemporaryDirectory(prefix="salonops-sop-") as directory:
        path = Path(directory) / "workbook.xlsx"
        path.write_bytes(payload)
        path.chmod(0o600)
        yield path


def run_import(payload, dry_run):
    output = io.StringIO()
    with private_workbook(payload) as path:
        try:
            call_command("import_sop_workbook", str(path), dry_run=dry_run, validate_workflow=True, stdout=output, no_color=True)
        except CommandError as exc:
            raise ImportProblem(str(exc)) from exc
        except (ObjectDoesNotExist, ValueError, TypeError, KeyError, DatabaseError) as exc:
            # Do not echo cell values, SQL, paths or workbook contents into pages/logs.
            raise ImportProblem("Import validation failed. Check identifiers, references, column values and field lengths.") from exc
    return output.getvalue().strip()


@transaction.atomic
def stage_workbook(user, filename, payload):
    validate_workbook(payload)
    lock_masters()
    before = master_snapshot()
    summary = run_import(payload, dry_run=True)
    # One pending workbook per administrator; no unattended repeated imports.
    SOPWorkbookImport.objects.filter(created_by=user, status="PENDING").update(
        status="CANCELLED", workbook=b"",
    )
    return SOPWorkbookImport.objects.create(
        created_by=user, filename=Path(filename.replace("\\", "/")).name[:160],
        sha256=hashlib.sha256(payload).hexdigest(), workbook=payload,
        master_digest=hashlib.sha256(before).hexdigest(), summary=summary,
        expires_at=timezone.now() + timedelta(hours=1),
    )


@transaction.atomic
def apply_workbook(stage_id, user):
    stage = SOPWorkbookImport.objects.select_for_update().get(pk=stage_id, created_by=user)
    if stage.status != "PENDING" or stage.expires_at <= timezone.now():
        raise ImportProblem("This preview expired or has already been used. Upload and validate again.")
    payload = bytes(stage.workbook)
    if hashlib.sha256(payload).hexdigest() != stage.sha256:
        raise ImportProblem("The staged workbook changed. Upload and validate again.")
    lock_masters()
    before = master_snapshot()
    if hashlib.sha256(before).hexdigest() != stage.master_digest:
        raise ImportProblem("Master data changed after validation. Upload again to review a fresh preview.")
    # Snapshot and master changes commit together or both roll back. No visit models touched.
    stage.master_snapshot = gzip.compress(before, mtime=0)
    stage.summary = run_import(payload, dry_run=False)
    stage.status = "APPLIED"
    stage.applied_at = timezone.now()
    stage.workbook = b""
    stage.save(update_fields=["master_snapshot", "summary", "status", "applied_at", "workbook"])
    return stage
