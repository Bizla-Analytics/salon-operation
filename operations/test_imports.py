from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory

from django.core.management import call_command
from django.test import TestCase
from openpyxl import Workbook

from .models import OperationalTask, Service, SubService


class WorkbookImportCompatibilityTests(TestCase):
    def workbook(self, directory, v3=False):
        # Minimal synthetic fixture, not customer or business master data.
        book = Workbook()
        book.remove(book.active)
        for name, headers in {
            'Service Master': ['service_code'],
            'Sub-Service Master': ['sub_service_code'],
            'Inventory Master': ['inventory_code'],
            'Equipment Master': ['equipment_code'],
            'Service Detail': ['service_detail_id'],
            'Task Inventory': ['task_inventory_id'],
            'Task Equipment': ['task_equipment_id'],
        }.items():
            book.create_sheet(name).append(headers)
        headers = ['task_id','sub_service_code','task_sequence','activity_code','task_name','role_code','employee_role','active_labour_minutes','passive_time_minutes','required_skill','allowed_staff_type','default_customer_type','is_skippable','notes','status']
        values = ['SYNTH_TASK','SYNTH_SUB',1,'ACT','Example','','',3,5,'','','',False,'Internal note','Active']
        if v3:
            headers += ['supports_waiting','staff_instructions']
            values += [True,'Apply, wait, resume, remove.']
        book.create_sheet('Task Master').append(headers)
        book['Task Master'].append(values)
        path = Path(directory) / 'fixture.xlsx'
        book.save(path)
        return str(path)

    def test_v2_remains_compatible_and_v3_dry_run_is_non_mutating(self):
        SubService.objects.create(code='SYNTH_SUB', name='Example')
        with TemporaryDirectory() as directory:
            path = self.workbook(directory)
            call_command('import_sop_workbook',path,stdout=StringIO())
            task = OperationalTask.objects.get(external_id='SYNTH_TASK')
            self.assertFalse(task.supports_waiting)
            self.assertEqual(task.staff_instructions,'')
            path = self.workbook(directory,v3=True)
            call_command('import_sop_workbook',path,dry_run=True,stdout=StringIO())
            task.refresh_from_db()
            self.assertFalse(task.supports_waiting)
            call_command('import_sop_workbook',path,stdout=StringIO())
            call_command('import_sop_workbook',path,stdout=StringIO())
            task.refresh_from_db()
            self.assertTrue(task.supports_waiting)
            self.assertEqual(task.staff_instructions,'Apply, wait, resume, remove.')
            self.assertEqual(OperationalTask.objects.count(),1)
